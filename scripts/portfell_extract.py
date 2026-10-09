import openpyxl, csv, sys, datetime
from openpyxl.utils import column_index_from_string as ci
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
def num(v):
    try: return float(v)
    except: return None
rows = []
def rd(x,n):
    if x is None: return ''
    return int(x) if n==0 else round(x,n)
def pct(x): return None if x is None else x*100
def dstr(v):
    return v.date().isoformat() if isinstance(v, datetime.datetime) else ''
for book, sheet in [('EU','EU AKTSIA'),('US','US AKTSIA'),('ASIA','AASIA AKTSIA')]:
    ws = wb[sheet]
    hdr_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), ())
    hdr = {str(v).strip(): i for i, v in enumerate(hdr_row) if v is not None}
    col = lambda name, letter: hdr.get(name, ci(letter) - 1)   # same lookup as channels.py
    c_cost, c_pl, c_real = col('Ost EUR','O'), col('€ Kasum','BE'), col('€ Rtud PL','BJ')
    c_price, c_date = col('Hind','I'), col('Ost Date','F')
    for r in ws.iter_rows(min_row=3, values_only=True):
        r = list(r) + [None]*80
        if r[0] and (num(r[24]) or 0) > 0:
            rows.append([book, r[0], r[1], r[2], r[3], rd(num(r[24]),0), rd(num(r[27]),4), rd(pct(num(r[29])),2), '', '', '',
                         rd(num(r[c_cost]),2), rd(num(r[c_pl]),2), rd(num(r[c_real]),2), rd(num(r[c_price]),6), dstr(r[c_date])])
for r in wb['BONDS'].iter_rows(min_row=5, values_only=True):
    r = list(r) + [None]*50
    nom = num(r[56])  # Saldo col 57 = open nominal
    if (r[0] or r[1]) and nom and nom > 0:
        rows.append(['BONDS', r[0], r[1], r[2], r[13], nom, rd(num(r[20]),3), rd(pct(num(r[22])),2), rd(pct(num(r[14])),3), rd(pct(num(r[35])),2), r[41] if not isinstance(r[41], datetime.datetime) else r[41].date().isoformat(),
                     '', '', '', '', ''])
with open(out, 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['Book','Symbol','Name','Country_or_Parent','Currency','Saldo_or_Nominal','Last','Daily_pct','Coupon_pct','Yield_pct','Maturity',
                'Cost_EUR','PL_EUR','Realised_EUR','Buy_price','Buy_date'])
    w.writerows(rows)
from collections import Counter
print(Counter(x[0] for x in rows))
